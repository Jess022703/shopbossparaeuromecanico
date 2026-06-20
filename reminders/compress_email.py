#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
compress_email.py
Sistema de reminders de mantenimiento — Euromecanico Corp
Taller especializado en Porsche · Sabana Seca / Toa Baja, Puerto Rico

Envia correos de recordatorio de mantenimiento (servicio cada ~6 meses) a los
clientes listados en un archivo Excel, usando Gmail por SMTP SSL.

Uso
---
1. Instalar dependencia:        python -m pip install openpyxl
2. Configurar la contrasena:    set GMAIL_APP_PASSWORD=xxxxxxxxxxxxxxxx   (Windows / PowerShell: $env:GMAIL_APP_PASSWORD="...")
   (Es una "App Password" de Google, NO la contrasena normal de la cuenta.)
3. (Opcional) Probar primero: poner TEST_MODE = True (envia solo a TEST_EMAIL,
   max 3). Ya viene en TEST_MODE = False para el envio real.
4. Ejecutar:                    python compress_email.py
   En modo real se pedira confirmacion ("escribe SI") antes de enviar.
5. Anti-reenvio: con EVITAR_REENVIO = True no se reenvia a un cliente que ya
   recibio el correo en los ultimos MESES_ESPERA meses (por defecto 6).

Las credenciales NUNCA se escriben en el codigo: la contrasena se lee de la
variable de entorno GMAIL_APP_PASSWORD.
"""

import os
import csv
import time
import calendar
import smtplib
from datetime import datetime
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

import openpyxl

# ---------------------------------------------------------------------------
# CONFIGURACION
# ---------------------------------------------------------------------------

# --- Modo prueba ---------------------------------------------------------
TEST_MODE = False                          # False = envia a los CLIENTES REALES
TEST_EMAIL = "admin@euromecanicocorp.com"  # destino unico cuando TEST_MODE = True
TEST_LIMIT = 3                             # cantidad maxima de correos de muestra

# --- Anti-reenvio (no volver a escribir al mismo cliente por un tiempo) ---
# True  = NO se reenvia a un cliente que ya recibio el correo en los ultimos
#         MESES_ESPERA meses (recomendado).
# False = se envia siempre, aunque ya se le haya escrito antes.
EVITAR_REENVIO = True
MESES_ESPERA = 6                       # ventana de espera antes de poder reenviar

# --- Cuenta de envio (Gmail) ---------------------------------------------
SENDER_EMAIL = "admin@euromecanicocorp.com"
SENDER_NAME = "Euromecanico Corp"
GMAIL_APP_PASSWORD = os.environ.get("GMAIL_APP_PASSWORD")  # NUNCA hardcodear
SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465

# --- Datos del taller (para el cuerpo del correo) ------------------------
SHOP_NAME = "EUROMECANICO CORP"
SHOP_PHONE = "(787) 344-6328"
SHOP_LOCATION = "Sabana Seca, Toa Baja, Puerto Rico"
BRAND_RED = "#E31A2B"

# --- Archivos ------------------------------------------------------------
EXCEL_FILE = "clientes_reminder_6_meses.xlsx"
EXCEL_SHEET = "Para Reminder 6 Meses"
# Log separado en modo prueba para no contaminar el log de produccion.
LOG_FILE = "compress_enviados_test.csv" if TEST_MODE else "compress_enviados.csv"

# --- Indices de columnas en el Excel -------------------------------------
COL_NOMBRE = 0      # Cliente CSV
COL_TELEFONO = 1    # Telefono
COL_EMAIL = 2       # Email
COL_VEHICULO = 3    # Vehiculo CSV (ej: "2013 PORSCHE CAYENNE TURBO")

# --- Pausa entre envios (segundos) ---------------------------------------
SEND_DELAY = 2

# Encabezado del log CSV
LOG_HEADER = ["email", "nombre", "vehiculo", "estado", "fecha"]


# ---------------------------------------------------------------------------
# UTILIDADES
# ---------------------------------------------------------------------------

def primer_nombre(nombre):
    """Devuelve el primer nombre en formato Title Case ('JORGE RIVERA' -> 'Jorge')."""
    if not nombre:
        return "Cliente"
    return nombre.strip().split()[0].title()


def formatear_vehiculo(vehiculo):
    """Pasa '2013 PORSCHE CAYENNE TURBO' -> '2013 Porsche Cayenne Turbo'."""
    if not vehiculo:
        return "su vehiculo"
    return " ".join(p.title() for p in str(vehiculo).strip().split())


def email_valido(email):
    """Validacion basica: tiene texto, una '@' y un '.' despues de la '@'."""
    if not email:
        return False
    email = str(email).strip()
    if email.count("@") != 1:
        return False
    local, _, dominio = email.partition("@")
    return bool(local) and "." in dominio


def hace_n_meses(fecha, meses):
    """Devuelve el datetime de 'meses' meses antes de 'fecha' (sin dependencias)."""
    anio = fecha.year
    mes = fecha.month - meses
    while mes <= 0:
        mes += 12
        anio -= 1
    dia = min(fecha.day, calendar.monthrange(anio, mes)[1])
    return fecha.replace(year=anio, month=mes, day=dia)


def cargar_historial(log_path):
    """Lee el log y devuelve {email: fecha_del_ultimo_envio_exitoso} (datetime).

    Solo cuentan los envios exitosos; asi el dedupe por ventana de tiempo no se
    confunde con intentos que terminaron en error.
    """
    historial = {}
    if not os.path.exists(log_path):
        return historial
    with open(log_path, "r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            correo = (row.get("email") or "").strip().lower()
            estado = (row.get("estado") or "").strip().lower()
            if not correo or not estado.startswith("enviado"):
                continue
            try:
                fecha = datetime.strptime((row.get("fecha") or "").strip(),
                                          "%Y-%m-%d %H:%M:%S")
            except ValueError:
                continue
            if correo not in historial or fecha > historial[correo]:
                historial[correo] = fecha
    return historial


def registrar_envio(log_path, email, nombre, vehiculo, estado):
    """Agrega una fila al log CSV (crea el encabezado si el archivo es nuevo)."""
    nuevo = not os.path.exists(log_path)
    with open(log_path, "a", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        if nuevo:
            writer.writerow(LOG_HEADER)
        writer.writerow([
            email,
            nombre,
            vehiculo,
            estado,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ])


# ---------------------------------------------------------------------------
# CONSTRUCCION DEL CORREO
# ---------------------------------------------------------------------------

def construir_html(nombre, vehiculo):
    """Email HTML profesional con header oscuro y marca en rojo."""
    saludo = primer_nombre(nombre)
    vehiculo_fmt = formatear_vehiculo(vehiculo)
    return f"""\
<!DOCTYPE html>
<html lang="es">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1.0"></head>
<body style="margin:0;padding:0;background-color:#1a1a1a;font-family:Arial,Helvetica,sans-serif;">
  <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background-color:#1a1a1a;padding:24px 0;">
    <tr><td align="center">
      <table role="presentation" width="600" cellpadding="0" cellspacing="0" style="max-width:600px;width:100%;background-color:#ffffff;border-radius:10px;overflow:hidden;">

        <!-- Header oscuro -->
        <tr>
          <td style="background-color:#0a0a0a;padding:28px 32px;text-align:center;border-bottom:3px solid {BRAND_RED};">
            <span style="color:{BRAND_RED};font-size:26px;font-weight:bold;letter-spacing:2px;">{SHOP_NAME}</span>
            <div style="color:#c9a227;font-size:11px;letter-spacing:3px;margin-top:6px;text-transform:uppercase;">Independent Porsche Specialists</div>
          </td>
        </tr>

        <!-- Cuerpo -->
        <tr>
          <td style="padding:32px;color:#222222;font-size:15px;line-height:1.6;">
            <p style="margin:0 0 16px;">Hola <strong>{saludo}</strong>,</p>

            <p style="margin:0 0 16px;">
              Le escribimos desde <strong>{SHOP_NAME}</strong> para recordarle que su
              <strong style="color:{BRAND_RED};">{vehiculo_fmt}</strong> esta proximo a cumplir
              su intervalo de mantenimiento recomendado (aproximadamente cada 6 meses).
            </p>

            <p style="margin:0 0 16px;">
              Mantener el servicio al dia ayuda a preservar el rendimiento, la fiabilidad y el
              valor de su Porsche. Nuestro equipo de especialistas estara encantado de coordinar
              una cita en el momento que mejor le convenga.
            </p>

            <!-- Llamada a la accion -->
            <table role="presentation" cellpadding="0" cellspacing="0" style="margin:24px 0;">
              <tr>
                <td style="background-color:{BRAND_RED};border-radius:6px;">
                  <a href="tel:+17873446328" style="display:inline-block;padding:14px 26px;color:#ffffff;font-size:16px;font-weight:bold;text-decoration:none;">
                    Llamar al {SHOP_PHONE}
                  </a>
                </td>
              </tr>
            </table>

            <p style="margin:0;color:#555555;font-size:14px;">
              Telefono del taller: <strong>{SHOP_PHONE}</strong><br>
              {SHOP_LOCATION}
            </p>
          </td>
        </tr>

        <!-- Footer -->
        <tr>
          <td style="background-color:#f2f2f2;padding:20px 32px;color:#888888;font-size:12px;line-height:1.5;text-align:center;">
            Si ya realizo este servicio recientemente o no desea recibir estos recordatorios,
            puede ignorar este mensaje.<br>
            &copy; {datetime.now().year} {SHOP_NAME} &middot; {SHOP_LOCATION}
          </td>
        </tr>

      </table>
    </td></tr>
  </table>
</body>
</html>"""


def construir_texto(nombre, vehiculo):
    """Version texto plano (fallback para clientes de correo sin HTML)."""
    saludo = primer_nombre(nombre)
    vehiculo_fmt = formatear_vehiculo(vehiculo)
    return (
        f"Hola {saludo},\n\n"
        f"Le escribimos desde {SHOP_NAME} para recordarle que su {vehiculo_fmt} esta "
        f"proximo a cumplir su intervalo de mantenimiento recomendado (cada ~6 meses).\n\n"
        f"Para coordinar una cita, llamenos al {SHOP_PHONE}.\n"
        f"{SHOP_LOCATION}\n\n"
        f"Si ya realizo este servicio o no desea recibir estos recordatorios, "
        f"puede ignorar este mensaje.\n"
    )


def construir_mensaje(destinatario, nombre, vehiculo):
    """Arma el MIMEMultipart (texto + HTML) listo para enviar."""
    msg = MIMEMultipart("alternative")
    msg["Subject"] = f"Recordatorio de mantenimiento - {formatear_vehiculo(vehiculo)}"
    msg["From"] = f"{SENDER_NAME} <{SENDER_EMAIL}>"
    msg["To"] = destinatario
    msg.attach(MIMEText(construir_texto(nombre, vehiculo), "plain", "utf-8"))
    msg.attach(MIMEText(construir_html(nombre, vehiculo), "html", "utf-8"))
    return msg


# ---------------------------------------------------------------------------
# LECTURA DEL EXCEL
# ---------------------------------------------------------------------------

def leer_clientes(ruta, hoja):
    """Lee el Excel y devuelve una lista de dicts {nombre, telefono, email, vehiculo}."""
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"No se encontro el archivo Excel: {ruta}")

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if hoja not in wb.sheetnames:
        raise ValueError(
            f"La hoja '{hoja}' no existe. Hojas disponibles: {wb.sheetnames}"
        )
    ws = wb[hoja]

    clientes = []
    # min_row=2 salta la fila de encabezados.
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or all(c is None for c in fila):
            continue

        def val(idx):
            return str(fila[idx]).strip() if idx < len(fila) and fila[idx] is not None else ""

        clientes.append({
            "nombre": val(COL_NOMBRE),
            "telefono": val(COL_TELEFONO),
            "email": val(COL_EMAIL),
            "vehiculo": val(COL_VEHICULO),
        })

    wb.close()
    return clientes


# ---------------------------------------------------------------------------
# PROGRAMA PRINCIPAL
# ---------------------------------------------------------------------------

def main():
    print("=" * 60)
    print(f"  {SHOP_NAME} - Reminders de mantenimiento (6 meses)")
    print(f"  Modo: {'PRUEBA' if TEST_MODE else 'PRODUCCION (envio real)'}")
    print("=" * 60)

    # 1) Validar credencial
    if not GMAIL_APP_PASSWORD:
        print("\nERROR: falta la variable de entorno GMAIL_APP_PASSWORD.")
        print('  Windows CMD : set GMAIL_APP_PASSWORD=tu_app_password')
        print('  PowerShell  : $env:GMAIL_APP_PASSWORD="tu_app_password"')
        return

    # 2) Leer Excel
    try:
        clientes = leer_clientes(EXCEL_FILE, EXCEL_SHEET)
    except (FileNotFoundError, ValueError) as e:
        print(f"\nERROR leyendo el Excel: {e}")
        return
    print(f"\nClientes en el Excel: {len(clientes)}")

    # 3) Cargar historial para no reenviar dentro de la ventana de espera
    historial = cargar_historial(LOG_FILE)
    cutoff = hace_n_meses(datetime.now(), MESES_ESPERA)
    if EVITAR_REENVIO:
        print(f"Anti-reenvio: ACTIVADO (no se reenvia si ya se envio en los "
              f"ultimos {MESES_ESPERA} meses)")
        if historial:
            print(f"Clientes con envio previo registrado ({LOG_FILE}): {len(historial)}")
    else:
        print("Anti-reenvio: DESACTIVADO (se enviara aunque ya se haya enviado antes)")

    # 4) Construir lista de trabajo (filtra emails invalidos y los que recibieron
    #    el correo dentro de la ventana de MESES_ESPERA meses)
    worklist = []
    omitidos_sin_email = 0
    omitidos_cooldown = 0
    for c in clientes:
        if not email_valido(c["email"]):
            omitidos_sin_email += 1
            continue
        if EVITAR_REENVIO:
            ultimo = historial.get(c["email"].strip().lower())
            if ultimo and ultimo >= cutoff:
                omitidos_cooldown += 1
                continue
        worklist.append(c)

    if TEST_MODE:
        worklist = worklist[:TEST_LIMIT]
        print(f"Modo prueba: se enviaran maximo {TEST_LIMIT} correos a {TEST_EMAIL}")

    print(f"Omitidos (email invalido/vacio):       {omitidos_sin_email}")
    print(f"Omitidos (enviado hace <{MESES_ESPERA} meses):  {omitidos_cooldown}")
    print(f"Por enviar:                            {len(worklist)}\n")

    if not worklist:
        print("No hay correos pendientes por enviar.")
        return

    # 4b) Confirmacion de seguridad antes de un envio REAL a clientes.
    if not TEST_MODE:
        print("=" * 60)
        print(f"  MODO REAL: se enviaran {len(worklist)} correos a CLIENTES REALES")
        print(f"  Remitente: {SENDER_NAME} <{SENDER_EMAIL}>")
        print("=" * 60)
        respuesta = input("Escribe 'SI' (en mayusculas) para confirmar el envio: ").strip()
        if respuesta != "SI":
            print("Cancelado. No se envio ningun correo.")
            return

    # 5) Enviar
    total = len(worklist)
    exitos = 0
    errores = 0

    try:
        servidor = smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT)
        servidor.login(SENDER_EMAIL, GMAIL_APP_PASSWORD)
    except smtplib.SMTPAuthenticationError:
        print("ERROR de autenticacion: revisa SENDER_EMAIL y GMAIL_APP_PASSWORD")
        print("  (debe ser una App Password de Google, no la contrasena normal).")
        return
    except Exception as e:
        print(f"ERROR conectando a {SMTP_HOST}:{SMTP_PORT} -> {e}")
        return

    try:
        for i, c in enumerate(worklist, start=1):
            destinatario = TEST_EMAIL if TEST_MODE else c["email"]
            try:
                mensaje = construir_mensaje(destinatario, c["nombre"], c["vehiculo"])
                servidor.sendmail(SENDER_EMAIL, destinatario, mensaje.as_string())
                estado = "enviado-prueba" if TEST_MODE else "enviado"
                exitos += 1
                print(f"[{i}/{total}] ✅ {c['nombre']} → {destinatario}")
            except Exception as e:
                estado = f"error: {e}"
                errores += 1
                print(f"[{i}/{total}] ❌ {c['nombre']} → {destinatario}  ({e})")

            # Log con el email REAL del cliente (asi el dedupe funciona por cliente).
            registrar_envio(LOG_FILE, c["email"], c["nombre"], c["vehiculo"], estado)

            # Pausa para no ser bloqueado por Gmail (salvo el ultimo).
            if i < total:
                time.sleep(SEND_DELAY)
    finally:
        servidor.quit()

    # 6) Resumen
    print("\n" + "=" * 60)
    print(f"  Enviados: {exitos}   Errores: {errores}   Total: {total}")
    print(f"  Log: {LOG_FILE}")
    print("=" * 60)


if __name__ == "__main__":
    main()
